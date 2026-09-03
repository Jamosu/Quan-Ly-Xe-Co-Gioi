import docx
import sys

sys.stdout.reconfigure(encoding='utf-8')

doc = docx.Document('BRD_QuanLyXeCoGioi_KLH.docx')

with open('BRD_extracted.txt', 'w', encoding='utf-8') as f:
    f.write(f"=== PARAGRAPHS ({len(doc.paragraphs)}) ===\n")
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip():
            f.write(f"[{i}] {p.text}\n")
            
    f.write(f"\n=== TABLES ({len(doc.tables)}) ===\n")
    for t_idx, table in enumerate(doc.tables):
        f.write(f"\n--- TABLE {t_idx+1} ({len(table.rows)} rows, {len(table.columns)} cols) ---\n")
        for row in table.rows:
            row_text = [cell.text.replace('\n', ' ').strip() for cell in row.cells]
            f.write(" | ".join(row_text) + "\n")

print("Extracted BRD successfully to BRD_extracted.txt")
