from copy import copy
from pathlib import Path

from openpyxl import load_workbook


TEMPLATE_DIR = Path(__file__).resolve().parent.parent / "FileTemplate_mau"
REFERENCE_FILE = TEMPLATE_DIR / "Template__CongTy.xlsx"


def copy_style(source, target):
    target.font = copy(source.font)
    target.fill = copy(source.fill)
    target.border = copy(source.border)
    target.alignment = copy(source.alignment)
    target.number_format = source.number_format
    target.protection = copy(source.protection)


reference_sheet = load_workbook(REFERENCE_FILE).active
title_style = reference_sheet["A1"]
header_style = reference_sheet["A3"]
instruction_style = reference_sheet["A4"]
data_style = reference_sheet["A5"]

for template_path in sorted(TEMPLATE_DIR.glob("Template_Import_*.xlsx")):
    workbook = load_workbook(template_path)
    for sheet_index, worksheet in enumerate(workbook.worksheets):
        if sheet_index == 0:
            copy_style(title_style, worksheet["A1"])
            worksheet.row_dimensions[1].height = 28
            worksheet.row_dimensions[2].height = 15
            for cell in worksheet[3]:
                copy_style(header_style, cell)
            worksheet.row_dimensions[3].height = 32
            for cell in worksheet[4]:
                copy_style(instruction_style, cell)
            worksheet.row_dimensions[4].height = 50
            for cell in worksheet[5]:
                copy_style(data_style, cell)
            worksheet.row_dimensions[5].height = 20
        else:
            for cell in worksheet[1]:
                copy_style(header_style, cell)
            worksheet.row_dimensions[1].height = 32

    workbook.save(template_path)

print(f"Styled operational templates from {REFERENCE_FILE.name}")
